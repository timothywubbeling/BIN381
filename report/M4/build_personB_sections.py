# build_personB_sections.py
# Builds Person B's Milestone 4 report sections (4, 6 and 8) as a .docx.
#
#   python report/M4/build_personB_sections.py
#
# Run from the repository root. Every number in the output is read from
# data/evaluation/ and data/powerbi/ at build time, so the report cannot drift
# away from the notebook. Re-run after any change to the model or the notebook.
#
# Section numbers follow the Milestone 4 team guideline:
#   4. Evaluation and key findings   (1.5 pages)
#   6. Recommendations               (1 page)
#   8. Limitations and future work   (0.5 pages)

import csv
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
EVAL = os.path.join(ROOT, "data", "evaluation")
PBI = os.path.join(ROOT, "data", "powerbi")
OUT = os.path.join(ROOT, "report", "M4", "BIN381_M4_PersonB_Sections_4_6_8.docx")


# ----------------------------------------------------------------------------
# Data
# ----------------------------------------------------------------------------
def read(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def num(row, key):
    return float(row[key])


ev = {r["model"]: r for r in read(os.path.join(EVAL, "evaluation_matrix.csv"))}
screen = read(os.path.join(EVAL, "screening_metrics_test.csv"))
sub = read(os.path.join(EVAL, "subgroup_performance.csv"))
imp = [r for r in read(os.path.join(EVAL, "permutation_importance.csv")) if r["model"] == "binary"]
orat = {r["term"]: r for r in read(os.path.join(EVAL, "odds_ratios.csv"))}
rates = read(os.path.join(EVAL, "observed_rates.csv"))
tier3 = {r["model"]: r for r in read(os.path.join(EVAL, "test_metrics_3tier.csv"))}
sev = {r["model"]: r for r in read(os.path.join(EVAL, "severe_flagged.csv"))}
wt = read(os.path.join(EVAL, "survey_weighted_check.csv"))[0]
diff = {r["comparison"]: r for r in read(os.path.join(EVAL, "bootstrap_diff_screen_auc.csv"))}
base = read(os.path.join(EVAL, "monitoring_baseline.csv"))[0]
recon = {(r["Breakdown"], r["Group"]): r for r in read(os.path.join(PBI, "powerbi_reconciliation.csv"))}

FINAL = "Binary logistic (glmnet)"
FOREST = "Random forest (ranger)"
f = ev[FINAL]

THR = num(f, "threshold")
RECALL = num(f, "screen_recall")
PREC = num(f, "screen_precision")
FLAG = num(f, "pct_flagged")
AUC = num(f, "screen_auc")
BRIER = num(f, "screen_brier")
ECE = num(f, "ece")
CV_AUC = num(f, "cv_screen_auc")
CV_GAP = num(f, "cv_test_gap")


def rule(model, rule_name):
    for r in screen:
        if r["model"] == model and r["rule"].startswith(rule_name):
            return r
    raise KeyError(rule_name)


f2_rule = rule(FINAL, "F2-optimal")
half_rule = rule(FINAL, "Default 0.5")


def sg(variable, group):
    for r in sub:
        if r["variable"] == variable and r["group"] == group:
            return r
    raise KeyError((variable, group))


def obs(variable, level):
    for r in rates:
        if r["variable"] == variable and r["level"] == level:
            return float(r["pct_food_insecure"])
    raise KeyError((variable, level))


# test-set size: every household appears once in the subgroup grant split
n_test = sum(int(r["households"]) for r in sub if r["variable"] == "grant_band")
n_at_risk = sum(int(r["at_risk"]) for r in sub if r["variable"] == "grant_band")
n_model = int(recon[("All", "All households")]["Households"])

imp.sort(key=lambda r: -float(r["auc_drop"]))

PRED_LABEL = {
    "ComparativeIncomeCode": "Income change vs a year ago",
    "Income_PerCapita": "Income per person",
    "SocialGrantRecipients": "Number of grant recipients",
    "Income_Log": "Household income (log)",
    "ElectricityAccessCode": "Electricity access",
    "GeoTypeCode": "Settlement type",
    "MainToiletCode": "Toilet type",
    "HouseholdSize": "Household size",
}


def pct(x, dp=1):
    return f"{100 * float(x):.{dp}f}%"


def wpct(breakdown, group):
    """Survey-weighted at-risk share from the dashboard reconciliation table."""
    return f"{float(recon[(breakdown, group)]['AtRiskPct_weighted']):.1f}%"


def od(term):
    r = orat[term]
    return f"{float(r['estimate']):.2f} ({float(r['conf.low']):.2f}–{float(r['conf.high']):.2f})"


# ----------------------------------------------------------------------------
# Document helpers
# ----------------------------------------------------------------------------
doc = Document()

# Match the Milestone 3 report exactly: A4, 2.54 cm margins. python-docx defaults to
# US Letter with 3.17 cm sides, which makes the page count here disagree with the page
# count once Timothy merges these sections into the real document.
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)

for name, size in (("Heading 1", 14), ("Heading 2", 12)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)


def para(text, style=None, italic=False, size=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    return p


def caption(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(2)
    return p


def bullet(text_bold, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text_bold).bold = True
    p.add_run(" " + text)
    p.paragraph_format.space_after = Pt(4)
    return p


def table(headers, rows, bold_rows=(), size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(size)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(size)
            run.bold = n in bold_rows
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# ----------------------------------------------------------------------------
# Section 4  -  Evaluation and key findings
# ----------------------------------------------------------------------------
doc.add_heading("4. Evaluation and key findings", level=1)
para("Rudi Jan du Plessis", italic=True, size=9)

para(
    f"All four candidates were tuned with stratified five-fold cross-validation on the training "
    f"set and then scored once on the untouched 30% test set ({n_test:,} households, "
    f"{n_at_risk:,} of them food insecure). Because the three-tier target could not separate the "
    f"thin middle tier — no three-tier model recalls more than "
    f"{pct(max(float(r['recall_insecure']) for r in tier3.values()))} of mildly-to-moderately "
    f"insecure households, and the best-ranking one manages "
    f"{pct(tier3[FOREST]['recall_insecure'])} — the "
    f"deployable question was reframed as a binary screen: is this household at risk of any food "
    f"insecurity? Models are therefore compared on how well they rank households by risk, not on "
    f"accuracy, which the {pct(1 - num(wt, 'unweighted_at_risk_rate'), 0)} food-secure majority "
    f"would flatter."
)

caption(
    "Table 4.1 Screening performance on the test set, each model at its own 80%-recall threshold "
    "(source: data/evaluation/evaluation_matrix.csv)"
)
table(
    ["Model", "ROC-AUC", "Threshold", "Recall", "Precision", "Referred", "Brier", "ECE"],
    [
        [
            m.replace(" (glmnet)", "").replace(" (ranger)", "").replace(" (rpart)", ""),
            f"{num(ev[m], 'screen_auc'):.3f}",
            f"{num(ev[m], 'threshold'):.2f}",
            pct(ev[m]["screen_recall"]),
            pct(ev[m]["screen_precision"]),
            pct(ev[m]["pct_flagged"]),
            f"{num(ev[m], 'screen_brier'):.3f}",
            f"{num(ev[m], 'ece'):.3f}",
        ]
        for m in [FINAL, FOREST, "Multinomial logistic (glmnet)", "Decision tree (rpart)"]
    ],
    bold_rows={0},  # the selected model
)

doc.add_heading("Why this model and this threshold", level=2)
para(
    f"The selected model is the lasso (glmnet) binary logistic regression at a referral threshold "
    f"of {THR:.2f}. The threshold was fixed on cross-validation, never on the test set, using a "
    f"policy rule agreed before evaluation: find at least 80% of at-risk households at the fewest "
    f"referrals. On the test set it finds {pct(RECALL)} of at-risk households and "
    f"{pct(sev[FINAL]['severe_flagged'])} of the severely food insecure, while referring "
    f"{pct(FLAG)} of households at {pct(PREC)} precision. The F2-optimal threshold of "
    f"{num(f2_rule, 'threshold'):.2f} was rejected because it refers "
    f"{pct(f2_rule['pct_flagged'])} of all households, and the default 0.5 because it misses "
    f"{pct(1 - num(half_rule, 'recall'))} of at-risk households. The random forest scores "
    f"marginally higher (AUC {num(ev[FOREST], 'screen_auc'):.3f} against {AUC:.3f}; bootstrap "
    f"difference +{float(diff['forest - binary']['mean_diff']):.3f}, 95% CI "
    f"{float(diff['forest - binary']['ci_low']):.3f}–"
    f"{float(diff['forest - binary']['ci_high']):.3f}), but the logistic model was chosen for one "
    f"interpretable odds ratio per factor, equally good calibration and a direct fit with the "
    f"Shiny app — a gain too small to pay for the loss of explanation in a tool that must "
    f"justify every referral to an official."
)
para(
    f"It generalises and it is calibrated: cross-validated AUC {CV_AUC:.3f} against {AUC:.3f} on "
    f"test (gap {CV_GAP:.4f}, so no meaningful overfitting), Brier {BRIER:.3f} and expected "
    f"calibration error {ECE:.3f}, so the scores are usable as probabilities and not only as a "
    f"ranking. Survey-weighting the test set moves recall only from "
    f"{pct(wt['unweighted_recall'])} to {pct(wt['weighted_recall'])}."
)

doc.add_heading("What drives risk, and what the data says about the hypotheses", level=2)
para(
    "Permutation importance ranks the drivers consistently across all three models: "
    + ", ".join(
        f"{PRED_LABEL[r['variable']].lower()} ({float(r['auc_drop']):.3f})" for r in imp[:4]
    )
    + ". Electricity, settlement type, toilet type and household size add little once income is "
    "known. The odds ratios below come from an unpenalised logistic fit on the training data "
    "(Person A's tuned penalty was effectively zero), so they describe the deployed model while "
    "adding confidence intervals — as associations, not causes."
)

caption("Table 4.2 Milestone 1 hypotheses against the final evidence (odds ratios with 95% CI)")
table(
    ["Hypothesis", "Evidence", "Verdict"],
    [
        [
            "H1 Falling income raises risk",
            f"{obs('ComparativeIncomeCode', '5'):.1f}% vs "
            f"{obs('ComparativeIncomeCode', '2'):.1f}–{obs('ComparativeIncomeCode', '3'):.1f}% "
            f"mid-scale; OR {od('ComparativeIncomeCode5')}. But code 1 (much higher) is also "
            f"{obs('ComparativeIncomeCode', '1'):.1f}%.",
            "Supported; U-shaped",
        ],
        [
            "H2 Rural households more at risk, income held equal",
            f"{obs('GeoTypeCode', 'Traditional_Tribal'):.1f}% vs "
            f"{obs('GeoTypeCode', 'Urban'):.1f}% urban raw; adjusted OR "
            f"{od('GeoTypeCodeTraditional_Tribal')}.",
            "Not supported; the gap is income and services",
        ],
        [
            "H3 Larger households more at risk at fixed income",
            f"OR {od('HouseholdSize')} per member; raw {obs('size', '1-2'):.1f}% → "
            f"{obs('size', '7+'):.1f}%, mostly via income per person.",
            "Supported; weak direct effect",
        ],
        [
            "H4 No electricity raises risk",
            f"{obs('ElectricityAccessCode', 'No'):.1f}% vs "
            f"{obs('ElectricityAccessCode', 'Yes'):.1f}%; OR {od('ElectricityAccessCodeNo')}. Only "
            f"~5% of households affected.",
            "Supported",
        ],
        [
            "H5 Grant receipt lowers risk",
            f"Reversed: {obs('grants', '1+ recipients'):.1f}% vs "
            f"{obs('grants', 'None'):.1f}%; OR {od('SocialGrantRecipients')} per recipient. Grants "
            f"are means-tested.",
            "Not supported; targeting, not cause",
        ],
        [
            "H6 Toilet type is a useful poverty proxy",
            f"Code 5 ({int([r for r in rates if r['variable'] == 'MainToiletCode' and r['level'] == '5'][0]['households']):,} households) "
            f"{obs('MainToiletCode', '5'):.1f}% vs {obs('MainToiletCode', '1'):.1f}%; OR "
            f"{od('MainToiletCode5')}. Importance near zero.",
            "Partly; little predictive value",
        ],
    ],
)

para(
    f"The gaps matter more than the averages. The screen finds "
    f"{pct(sg('income_band', '< R2,500')['recall'])} of at-risk households earning under R2,500 a "
    f"month but only {pct(sg('income_band', '> R11,500')['recall'])} of those over R11,500, and "
    f"{pct(sg('grant_band', 'No grants')['recall'])} of households with no grant recipients "
    f"against {pct(sg('grant_band', 'Receives grants')['recall'])} of those with them. By province "
    f"recall runs from {pct(sg('province', 'Gauteng')['recall'])} (Gauteng) to "
    f"{pct(sg('province', 'Eastern Cape')['recall'])} (Eastern Cape). It is weakest exactly where "
    f"food insecurity is least expected — a reason to keep self-referral open rather than "
    f"trust a ‘lower risk’ result."
)

# No page break between sections: these drop into the middle of the merged report,
# where Sections 5 and 7 (Person C) sit either side of Section 6.

# ----------------------------------------------------------------------------
# Section 6  -  Recommendations
# ----------------------------------------------------------------------------
doc.add_heading("6. Recommendations", level=1)
para("Rudi Jan du Plessis", italic=True, size=9)

para(
    "Each recommendation below starts from a number in Section 4 and names who should act. All of "
    "them assume the screening conditions in Section 5: an official reviews every referral, and "
    "the tool never decides eligibility on its own."
)

bullet(
    "1. Deploy the screen for outreach triage only, and resource it for the referral rate it "
    "actually produces.",
    f"At threshold {THR:.2f} the tool refers {pct(FLAG)} of households and "
    f"{pct(1 - PREC)} of those referrals turn out to be food secure. Programme managers should "
    f"size follow-up capacity against that rate before go-live, or raise the threshold knowingly "
    f"and accept lower recall; it is a policy choice, not a technical one. Monitoring should raise "
    f"an alert if recall on followed-up households falls below "
    f"{pct(base['recall_alert_below'])}.",
)
bullet(
    "2. Ask every household about income change over the past year at intake.",
    f"It is the single strongest predictor (AUC drop {float(imp[0]['auc_drop']):.3f}, roughly "
    f"{float(imp[0]['auc_drop']) / float(imp[1]['auc_drop']):.1f}× the next variable), and "
    f"{obs('ComparativeIncomeCode', '5'):.1f}% of households reporting a sharp fall are food "
    f"insecure. Field staff can collect it in one question, and it should trigger follow-up on its "
    f"own even before the full screen is run.",
)
bullet(
    "3. Treat grant-receiving households as a group to check, not a group to deprioritise.",
    f"{obs('grants', '1+ recipients'):.1f}% of households with grant recipients are food insecure "
    f"against {obs('grants', 'None'):.1f}% without. Social development should coordinate "
    f"food-relief outreach with existing grant registers, which is the cheapest route to the "
    f"households most at risk — while reading the association as evidence of targeting, not of "
    f"grants causing harm.",
)
bullet(
    "4. Direct food-relief and basic-services budgets to the four provinces with the highest "
    "weighted risk first.",
    f"Weighted to the population, "
    f"{wpct('Province', 'Eastern Cape')} of Eastern Cape households are "
    f"at risk, {wpct('Province', 'Northern Cape')} in the Northern Cape, "
    f"{wpct('Province', 'North West')} in North West and "
    f"{wpct('Province', 'Mpumalanga')} in Mpumalanga, against "
    f"{wpct('All', 'All households')} nationally. Provincial teams can "
    f"drill into settlement type, income band and grant status within each province on the "
    f"dashboard.",
)
bullet(
    "5. Run a standing audit of ‘lower risk’ households in the groups where the model is "
    "weakest.",
    f"Recall falls to {pct(sg('income_band', '> R11,500')['recall'])} for households earning over "
    f"R11,500 and {pct(sg('grant_band', 'No grants')['recall'])} for households receiving no "
    f"grants. A monthly sample of lower-risk households in those groups, reviewed by field "
    f"officers, is the only way the institute will see the cases the screen misses. Keep "
    f"self-referral open and never present a low score as clearance.",
)

# ----------------------------------------------------------------------------
# Section 8  -  Limitations and future work
# ----------------------------------------------------------------------------
doc.add_heading("8. Limitations and future work", level=1)
para("Rudi Jan du Plessis", italic=True, size=9)

para(
    f"Modest discrimination is the central limitation. An AUC of {AUC:.3f} means the screen "
    f"separates risk only moderately: at the chosen threshold {pct(1 - PREC)} of referred "
    f"households are not food insecure, and {pct(1 - RECALL)} of at-risk households are missed. "
    f"The eight predictors describe a household's resources, not its access to food, and the "
    f"strongest of them is a single self-reported question about income change."
)
para(
    f"The evidence is one cross-sectional wave of the General Household Survey 2024, so every "
    f"result is an association rather than a cause; the reversed grant finding (H5) is the clearest "
    f"case where targeting and effect cannot be separated. The target is itself a proxy built from "
    f"self-reported items, which carry recall and social-desirability bias, and its three-tier "
    f"version was abandoned because the middle tier is not separable with these predictors, so the "
    f"tool cannot grade severity. Training used the unweighted sample (survey-weighting moves "
    f"recall by {abs(num(wt, 'unweighted_recall') - num(wt, 'weighted_recall')) * 100:.1f} points); "
    f"the person-level files D01–D03 and the D11 labour survey were excluded in Milestone 2; "
    f"and small subgroups such as farms "
    f"({int(sg('GeoTypeCode', 'Farms')['households']):,} test households) have intervals too wide "
    f"to judge. Most importantly, the model has never been validated outside this survey: there is "
    f"no external cohort and no record of what officials found when they followed up."
)
para(
    "Future work follows from those gaps, in order of value. First, capture follow-up outcomes so "
    "the next model can be trained and validated on what was actually found rather than on a "
    "survey proxy. Second, add a second GHS wave to allow within-household change and a genuine "
    "test of whether grants protect against food insecurity. Third, bring in per-person income, "
    "employment, debt and local food prices, which the current eight predictors stand in for only "
    "crudely. Fourth, revisit severity grading once a richer target exists, and re-tune the "
    "threshold against the institute's measured follow-up capacity instead of an assumed 80% "
    "recall floor."
)

# ----------------------------------------------------------------------------
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("wrote", os.path.relpath(OUT, ROOT))
print(f"  final model   : {FINAL} @ {THR:.2f}")
print(f"  test set      : {n_test:,} households, {n_at_risk:,} at risk")
print(f"  modelling set : {n_model:,} households")
print(f"  recall/prec   : {pct(RECALL)} / {pct(PREC)}, {pct(FLAG)} referred, AUC {AUC:.3f}")
