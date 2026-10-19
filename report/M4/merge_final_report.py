#!/usr/bin/env python3
"""Build the final Milestone 4 report from the three parts.

Alexander compiled sections 1 to 4, 6 and 8. This inserts Person C's sections 5 and 7 in
their numbered places and appends the ethics and AI log as Appendix B. Section numbers are
not written anywhere: they come from the report's own multilevel list, so inserting each
section in the right place is what renumbers them (Recommendations 5 -> 6, Limitations 6 -> 8).

Run from the repository root:  python report/M4/merge_final_report.py
"""
import argparse
import copy
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

HERE = os.path.dirname(os.path.abspath(__file__))

ap = argparse.ArgumentParser()
ap.add_argument("--base", default=os.path.join(HERE, "BIN381_M4_Report_base_PersonA.docx"))
ap.add_argument("--person-c", default=os.path.join(HERE, "BIN381_M4_PersonC_Sections_5_7.docx"))
ap.add_argument("--ethics", default=os.path.join(HERE, "BIN381_M4_Ethics_and_AI_Usage_Log.docx"))
ap.add_argument("--out", default=os.path.join(HERE, "BIN381_M4_Report_Group_P1.docx"))
args = ap.parse_args()

W_P, W_TBL, W_SECTPR, W_PPR = qn("w:p"), qn("w:tbl"), qn("w:sectPr"), qn("w:pPr")

# Captions are renumbered into one sequence across the merged report. Appendix A's table
# moves first, before Person C's tables are allowed to claim "Table 6".
APPENDIX_A_CAPTION = [("Table 6  Use of AI tools", "Table 9  Use of AI tools")]
PERSON_C_CAPTIONS = [("Table 5.1", "Table 6"), ("Table 5.2", "Table 7"),
                     ("Table 7.1", "Table 8"), ("Figures 5.1 and 5.2", "Figures 3 and 4"),
                     ("Figures 5.1-5.2", "Figures 3-4")]
ETHICS_CAPTIONS = [("Table 1.", "Table 10."), ("Table 2.", "Table 11.")]

# Notes to the team that must not reach the marker. They are reported on stdout instead.
ETHICS_DROP = ("Before upload:", "Draft for Timothy to rewrite")


def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def top_level(doc):
    """Body children only: paragraphs inside tables and inside the contents must not match."""
    return [el for el in doc.element.body if el.tag != W_SECTPR]


def anchor(doc, prefix):
    hits = [el for el in top_level(doc) if el.tag == W_P and text_of(el).startswith(prefix)]
    assert len(hits) == 1, "expected one paragraph starting %r, found %d" % (prefix, len(hits))
    return hits[0]


def block(doc, start, stop=None):
    """Copy of the body elements from the paragraph starting `start` up to (not including) `stop`."""
    kids = top_level(doc)
    first = next(i for i, el in enumerate(kids) if el.tag == W_P and text_of(el).startswith(start))
    last = len(kids)
    if stop:
        last = next(i for i, el in enumerate(kids) if el.tag == W_P and text_of(el).startswith(stop))
    return [copy.deepcopy(el) for el in kids[first:last]]


def style_of(el):
    pPr = el.find(W_PPR)
    if pPr is None:
        return None
    st = pPr.find(qn("w:pStyle"))
    return st.get(qn("w:val")) if st is not None else None


def replace_text(els, pairs):
    """Replace in run text, and insist every pair matches at least once."""
    for old, new in pairs:
        hits = 0
        for el in els:
            for t in el.iter(qn("w:t")):
                if t.text and old in t.text:
                    hits += t.text.count(old)
                    t.text = t.text.replace(old, new)
        assert hits, "caption %r not found, so it was never renumbered" % old


base = Document(args.base)
body = base.element.body
sectPr = body.find(W_SECTPR)

# Formatting templates taken from the report itself, so inserted content cannot drift from it.
TEMPLATE = {
    "section": copy.deepcopy(anchor(base, "Evaluation and key findings").find(W_PPR)),
    "sub": copy.deepcopy(anchor(base, "Target, predictors and data split").find(W_PPR)),
    "sub_plain": copy.deepcopy(anchor(base, "4.1 Why this model").find(W_PPR)),
    "appendix": copy.deepcopy(anchor(base, "Appendix A.").find(W_PPR)),
    "body": copy.deepcopy(anchor(base, "Households that cannot").find(W_PPR)),
    "bullet": copy.deepcopy(anchor(base, "BO1.").find(W_PPR)),
}


def restyle(el, kind, strip=None):
    pPr = el.find(W_PPR)
    if pPr is not None:
        el.remove(pPr)
    el.insert(0, copy.deepcopy(TEMPLATE[kind]))
    if not strip:
        return
    t = el.find(qn("w:r")).find(qn("w:t"))
    # Section headings must carry a number for the list to take over; section 7's
    # sub-headings were written without one, and the list numbers them anyway.
    assert kind != "section" or re.match(strip, t.text or ""), \
        "section heading %r has no number for the list to replace" % text_of(el)
    t.text = re.sub(strip, "", t.text or "", count=1)


def prepare(els, headings="numbered"):
    """Map the source document's styles onto the report's own."""
    for el in els:
        if el.tag != W_P:
            continue
        style = style_of(el)
        if style == "Heading1" and headings == "numbered":
            restyle(el, "section", strip=r"^\d+\.\s*")     # "5. Dashboard" -> the list numbers it
        elif style == "Heading1":
            restyle(el, "sub_plain")                       # appendix: keep the literal "1. Purpose"
        elif style == "Heading2" and headings == "numbered":
            restyle(el, "sub", strip=r"^\d+\.\d+\s*")      # "5.1 The dashboard" -> the list numbers it
        elif style == "Heading2":
            restyle(el, "sub_plain")
        elif style in ("ListBullet", "ListNumber"):
            restyle(el, "bullet")
        elif style is None and el.find(W_PPR) is None:
            el.insert(0, copy.deepcopy(TEMPLATE["body"]))  # match the spacing of the text around it
    return els


# --- Appendix A's table caption moves out of the way first -------------------------------
replace_text(top_level(base), APPENDIX_A_CAPTION)

# --- Person C's sections 5 and 7 ---------------------------------------------------------
person_c = Document(args.person_c)
section5 = prepare(block(person_c, "5. Dashboard and deployment", "7. Ethics and responsible AI"))
section7 = prepare(block(person_c, "7. Ethics and responsible AI"))
replace_text(section5 + section7, PERSON_C_CAPTIONS)

before_recommendations = anchor(base, "Recommendations")
for el in section5:
    before_recommendations.addprevious(el)

# --- The section 5 screenshots --------------------------------------------------------
# Person C's section carries a bracketed placeholder instead of the dashboard figures. The
# pictures have to be added here, to the finished document, so their relationships belong to
# it. The merged dashboard's own exports are preferred and picked up as soon as they exist.
FIGURES = os.path.abspath(os.path.join(HERE, "..", "..", "docs", "figures"))
WANTED = [("pbi_m4_1_overview.png", "the landing page"),
          ("pbi_m4_4_performance.png", "test-set performance with the threshold what-if")]
FALLBACK = {"pbi_m4_4_performance.png": "pbi_b1_performance.png"}


def available():
    for name, what in WANTED:
        path = os.path.join(FIGURES, name)
        if os.path.exists(path):
            yield path, what, False
        elif name in FALLBACK and os.path.exists(os.path.join(FIGURES, FALLBACK[name])):
            yield os.path.join(FIGURES, FALLBACK[name]), what, True


placeholder = anchor(base, "[Figures")
shots = list(available())
substitutes = [os.path.basename(p) for p, _, sub in shots if sub]
if shots:
    width = Cm(15.5 / len(shots))
    holder = base.add_paragraph()                      # built at the end, then moved into place
    holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, (path, _, _) in enumerate(shots):
        if i:
            holder.add_run("  ")
        holder.add_run().add_picture(path, width=width)
    placeholder.addprevious(holder._p)
    caption = anchor(base, "Figures 3-4")
    if len(shots) == 1:
        caption.find(qn("w:r")).find(qn("w:t")).text = (
            "Figure 3  " + shots[0][1][0].upper() + shots[0][1][1:] + ".")
    placeholder.getparent().remove(placeholder)
else:
    substitutes = ["no dashboard screenshot found; the placeholder is still in section 5"]
before_limitations = anchor(base, "Limitations and future work")
for el in section7:
    before_limitations.addprevious(el)

# --- The ethics and AI log as Appendix B -------------------------------------------------
ethics = Document(args.ethics)
scope = [el for el in top_level(ethics) if el.tag == W_P
         and text_of(el).startswith(("Ethics decisions and use of AI", "BIN381 SDHSA project"))]
assert len(scope) == 2, "expected the log's two scope paragraphs, found %d" % len(scope)
log = [copy.deepcopy(el) for el in scope] + block(ethics, "1. Purpose")

dropped = [text_of(el) for el in log if text_of(el).startswith(ETHICS_DROP)]
log = [el for el in log if not text_of(el).startswith(ETHICS_DROP)]
prepare(log, headings="plain")
replace_text(log, ETHICS_CAPTIONS)

heading_b = copy.deepcopy(anchor(base, "Appendix A."))
for extra in heading_b.findall(qn("w:bookmarkStart")) + heading_b.findall(qn("w:bookmarkEnd")):
    heading_b.remove(extra)
for run in heading_b.findall(qn("w:r"))[1:]:
    heading_b.remove(run)
heading_b.find(qn("w:r")).find(qn("w:t")).text = "Appendix B. Ethics report and AI usage log"
sectPr.addprevious(heading_b)
for el in log:
    sectPr.addprevious(el)

# --- One font for the body text, since the generated sections carry none -----------------
# Sections 1 to 3 are explicitly Calibri; without this, everything generated from a script
# renders in the theme font (Aptos) and the report looks like three different documents.
fonts = 0
for el in body:
    if el.tag not in (W_P, W_TBL):          # leave the cover and the contents control alone
        continue
    for p in ([el] if el.tag == W_P else el.iter(W_P)):
        if (style_of(p) or "").startswith(("Heading", "Title", "TOC")):
            continue
        for run in p.iter(qn("w:r")):
            rPr = run.find(qn("w:rPr"))
            if rPr is None:
                rPr = run.makeelement(qn("w:rPr"), {})
                run.insert(0, rPr)
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = rPr.makeelement(qn("w:rFonts"), {})
                rPr.insert(0, rFonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rFonts.set(qn(attr), "Calibri")
            fonts += 1

# --- Make Word rebuild the table of contents on open -------------------------------------
toc = next(el for el in body if any("TOC" in (i.text or "") for i in el.iter(qn("w:instrText"))))
begin = next(f for f in toc.iter(qn("w:fldChar")) if f.get(qn("w:fldCharType")) == "begin")
begin.set(qn("w:dirty"), "true")

base.save(args.out)

# --- Checks ------------------------------------------------------------------------------
out = Document(args.out)
kids = top_level(out)


def headings(numbered):
    found = []
    for el in kids:
        if el.tag != W_P or style_of(el) != "Heading2":
            continue
        if (el.find(W_PPR).find(qn("w:numPr")) is not None) == numbered:
            found.append(text_of(el))
    return found


assert headings(True) == [
    "Business problem and approach", "Data strategy, integration and preparation", "Modelling",
    "Evaluation and key findings", "Dashboard and deployment", "Recommendations",
    "Ethics and responsible AI", "Limitations and future work",
], headings(True)
assert headings(False) == ["References", "Appendix A. AI usage log, Rudi Jan du Plessis (Person B)",
                           "Appendix B. Ethics report and AI usage log"], headings(False)

full = "\n".join(text_of(el) for el in kids)
for stale in ("Table 5.1", "Table 5.2", "Table 7.1", "Figures 5.1", "Figure 5.1"):
    assert stale not in full, "%r survived the renumbering" % stale
for note in ETHICS_DROP:
    assert note not in full, "team note %r is still in the report" % note
def caption_numbers(word):
    """The numbers on the caption paragraphs, in the order they appear."""
    return [int(m.group(1)) for m in
            (re.match(word + r"s? (\d+)", text_of(el)) for el in kids if el.tag == W_P) if m]


assert caption_numbers("Table") == list(range(1, 12)), caption_numbers("Table")
assert caption_numbers("Figure") == [1, 2, 3], caption_numbers("Figure")
assert len(out.tables) == 11, len(out.tables)

print("wrote", args.out)
print("  sections  : " + " | ".join("%d. %s" % (i, h) for i, h in enumerate(headings(True), 1)))
print("  appendices: " + " | ".join(headings(False)[1:]))
print("  tables    : %d, captions renumbered to Table 1-11" % len(out.tables))
print("  font      : Calibri forced on %d body runs" % fonts)
print("  contents  : marked dirty, so Word rebuilds it when the file is opened")
for name in substitutes:
    print("  figure     : used " + name + " until the merged dashboard is exported")
for note in dropped:
    print("  dropped team note: " + note[:90])
